"""Genera la presentación final de MLOps II."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0F, 0x2C, 0x3C)
TEAL = RGBColor(0x1F, 0x6F, 0x78)
INK = RGBColor(0x1C, 0x28, 0x30)
MUTED = RGBColor(0x5C, 0x6B, 0x73)
GOLD = RGBColor(0xB0, 0x89, 0x48)
SAND = RGBColor(0xF4, 0xF1, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xDC, 0xD2)
SOFT = RGBColor(0xE7, 0xF2, 0xF2)
CREAM = RGBColor(0xFB, 0xF9, 0xF5)

W = Inches(13.333)
H = Inches(7.5)


def rgb(color):
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_run(run, text, size, color, bold=False, font="Calibri"):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False


def add_rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_round(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    # tighter corner
    shape.adjustments[0] = 0.08
    return shape


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    set_run(p.add_run(), text, size, color, bold, font)
    return box


def write_lines(slide, lines, x, y, w, h, size=15, color=INK, bold=False, spacing=6, align=PP_ALIGN.LEFT):
    """lines: list of str or (text, size, color, bold)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(spacing)
        if isinstance(line, tuple):
            text, sz, col, bd = line
        else:
            text, sz, col, bd = line, size, color, bold
        set_run(p.add_run(), text, sz, col, bd)
    return box


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, page, total=7):
    add_rect(slide, Inches(0), Inches(7.22), W, Inches(0.28), NAVY)
    add_text(slide, "MLOps II  ·  Satisfacción de pasajeros  ·  Belén Della Casa", Inches(0.5), Inches(7.25), Inches(9.5), Inches(0.22), 11, WHITE, False, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, f"{page}  /  {total}", Inches(11.4), Inches(7.25), Inches(1.45), Inches(0.22), 11, WHITE, False, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, title, kicker=None):
    add_rect(slide, Inches(0), Inches(0), W, Inches(0.08), TEAL)
    add_text(slide, title, Inches(0.5), Inches(0.24), Inches(12.2), Inches(0.48), 28, NAVY, True)
    add_rect(slide, Inches(0.52), Inches(0.78), Inches(1.15), Inches(0.045), GOLD)
    if kicker:
        add_text(slide, kicker, Inches(0.5), Inches(0.90), Inches(12.3), Inches(0.36), 15, MUTED, False)


def style_table(table, header=True, font=13, alignments=None):
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            solid = tcPr.find(qn("a:solidFill"))
            # fill via cell.fill
            cell.fill.solid()
            if r == 0 and header:
                cell.fill.fore_color.rgb = NAVY
                color = WHITE
                bold = True
            elif r % 2 == 0:
                cell.fill.fore_color.rgb = SAND
                color = INK
                bold = False
            else:
                cell.fill.fore_color.rgb = WHITE
                color = INK
                bold = False
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if not alignments else alignments[c]
            if p.runs:
                set_run(p.runs[0], p.runs[0].text, font, color, bold)
            else:
                set_run(p.add_run(), "", font, color, bold)
            # kill default borders a bit by setting them to LINE
            for border in ("lnL", "lnR", "lnT", "lnB"):
                ln = tcPr.find(qn(f"a:{border}"))
                if ln is None:
                    ln = tcPr.makeelement(qn(f"a:{border}"), {"w": "6350", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
                    tcPr.append(ln)
                else:
                    ln.set("w", "6350")
                # solid line color
                for child in list(ln):
                    ln.remove(child)
                sf = ln.makeelement(qn("a:solidFill"), {})
                srgb = sf.makeelement(qn("a:srgbClr"), {"val": "E2DCD2"})
                sf.append(srgb)
                ln.append(sf)


def fill_cell(cell, text):
    cell.text = text


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Predicción de satisfacción de pasajeros"
    prs.core_properties.author = "Belén Della Casa"
    prs.core_properties.subject = "MLOps II"
    blank = prs.slide_layouts[6]

    # ----- 1. Título -----
    s = prs.slides.add_slide(blank)
    set_bg(s, NAVY)
    add_rect(s, Inches(0), Inches(0), Inches(0.18), H, TEAL)
    add_text(s, "CEIA  ·  FIUBA  ·  OPERACIONES DE APRENDIZAJE AUTOMÁTICO II", Inches(0.7), Inches(1.35), Inches(11.5), Inches(0.32), 13, GOLD, True)
    add_text(s, "Predicción de satisfacción", Inches(0.7), Inches(1.85), Inches(11.8), Inches(0.62), 40, WHITE, True)
    add_text(s, "de pasajeros", Inches(0.7), Inches(2.5), Inches(11.8), Inches(0.62), 40, WHITE, True)
    add_rect(s, Inches(0.72), Inches(3.28), Inches(1.4), Inches(0.045), GOLD)
    add_text(s, "Plataforma MLOps para inferencia en tiempo real", Inches(0.7), Inches(3.5), Inches(11), Inches(0.4), 20, RGBColor(0xD5, 0xE4, 0xE6), False)
    add_text(s, "Belén Della Casa", Inches(0.7), Inches(5.55), Inches(8), Inches(0.35), 18, WHITE, True)
    notes(s, "Abrir con el caso en una frase: predecir si el pasajero queda satisfecho usando solo datos que la aerolínea ya tiene al cierre del vuelo. La presentación describe el servicio y cierra con la demo. Si MLflow está abierto, el ROC AUC se muestra ahí, no se cita de memoria.")

    # ----- 2. El caso -----
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    header(s, "El caso")
    add_text(s, "Anticipar la satisfacción del pasajero con lo que la aerolínea ya conoce al cierre del vuelo.", Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.4), 16, MUTED)
    points = [
        ("Clasificación binaria", "satisfied frente a neutral or dissatisfied."),
        ("Ocho features operativas", "Perfil del pasajero, clase, distancia y demoras. Sin puntajes de encuesta: esos puntajes casi son la respuesta."),
        ("Dos modelos", "XGBoost y Random Forest. El champion es el de mayor ROC AUC en test. El segundo queda como challenger."),
    ]
    y = 1.55
    for title, body in points:
        add_round(s, Inches(0.5), Inches(y), Inches(7.55), Inches(1.15), WHITE, LINE)
        add_rect(s, Inches(0.5), Inches(y), Inches(0.08), Inches(1.15), TEAL)
        add_text(s, title, Inches(0.8), Inches(y + 0.14), Inches(7.0), Inches(0.32), 16, NAVY, True)
        add_text(s, body, Inches(0.8), Inches(y + 0.48), Inches(7.0), Inches(0.55), 14, INK)
        y += 1.28

    add_round(s, Inches(8.3), Inches(1.55), Inches(4.5), Inches(4.85), NAVY)
    add_text(s, "FEATURES DEL CHAMPION", Inches(8.55), Inches(1.75), Inches(4.05), Inches(0.28), 12, GOLD, True)
    add_text(s, "airline_satisfaction_without_scores", Inches(8.55), Inches(2.08), Inches(4.05), Inches(0.5), 14, WHITE, True)
    feats = [
        "Gender",
        "Customer Type",
        "Type of Travel",
        "Class",
        "Age",
        "Flight Distance",
        "Departure Delay in Minutes",
        "Arrival Delay in Minutes",
    ]
    write_lines(s, feats, Inches(8.55), Inches(2.7), Inches(4.0), Inches(3.3), size=15, color=RGBColor(0xE7, 0xF2, 0xF2), spacing=4)
    footer(s, 2)
    notes(s, "Remarcar por qué no se usan los scores de la encuesta: el servicio tiene que funcionar con datos operativos. El nombre del modelo en el registry es airline_satisfaction_without_scores. No citar un AUC si no está a la vista en MLflow.")

    # ----- 3. De la notebook al servicio -----
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    header(s, "La plataforma", "Un solo servicio: entrena, registra el champion y lo expone para predecir.")
    rows = [
        ("Airflow", "Reentrenamiento semanal, domingos 23:59, o manual"),
        ("MLflow + PostgreSQL", "Experimentos, métricas y Model Registry"),
        ("MinIO", "Datasets y artefactos, sin depender de Drive en cada entrenamiento"),
        ("FastAPI", "Predicción online del champion"),
        ("Streamlit", "La misma predicción, para una persona"),
        ("Docker Compose", "Todo el entorno se levanta junto"),
    ]
    y = 1.45
    for i, (name, desc) in enumerate(rows):
        bg = WHITE if i % 2 == 0 else SAND
        add_round(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.82), bg, LINE)
        add_text(s, name, Inches(0.75), Inches(y + 0.22), Inches(3.3), Inches(0.4), 16, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, desc, Inches(4.15), Inches(y + 0.22), Inches(8.3), Inches(0.4), 16, INK, False, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.9
    footer(s, 3)
    notes(s, "Recorrer la plataforma como un solo sistema. Airflow orquesta, MLflow versiona, MinIO guarda datos y artefactos, FastAPI sirve el champion y Streamlit es la misma predicción para una persona. Docker Compose levanta todo junto.")

    # ----- 4. Arquitectura -----
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    header(s, "Arquitectura", "El entrenamiento y la predicción corren separados.")
    add_round(s, Inches(0.45), Inches(1.45), Inches(6.15), Inches(4.15), WHITE, LINE)
    add_rect(s, Inches(0.45), Inches(1.45), Inches(6.15), Inches(0.48), TEAL)
    add_text(s, "Entrenamiento", Inches(0.65), Inches(1.52), Inches(5.7), Inches(0.36), 16, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    train = [
        "1   Airflow dispara el pipeline",
        "2   Valida el CSV. Si falla, no entrena",
        "3   Entrena, evalúa y registra en MLflow",
        "4   Reasigna champion y challenger por ROC AUC",
        "5   POST /reload-model actualiza la API sin reiniciar",
    ]
    write_lines(s, train, Inches(0.7), Inches(2.15), Inches(5.65), Inches(3.2), size=15, color=INK, spacing=10)

    add_round(s, Inches(6.8), Inches(1.45), Inches(6.05), Inches(4.15), WHITE, LINE)
    add_rect(s, Inches(6.8), Inches(1.45), Inches(6.05), Inches(0.48), NAVY)
    add_text(s, "Predicción", Inches(7.0), Inches(1.52), Inches(5.6), Inches(0.36), 16, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    pred = [
        "1   Streamlit o un cliente llaman a POST /predict",
        "2   Pydantic valida la entrada",
        "3   El champion responde desde memoria",
        "4   La respuesta trae clase, probabilidad y versión",
    ]
    write_lines(s, pred, Inches(7.05), Inches(2.15), Inches(5.55), Inches(3.2), size=15, color=INK, spacing=12)

    urls = [
        ("Streamlit", ":8501"),
        ("API", ":8000"),
        ("MLflow", ":5001"),
        ("Airflow", ":8080"),
        ("MinIO", ":9101"),
    ]
    x = 0.45
    for name, port in urls:
        add_round(s, Inches(x), Inches(5.78), Inches(2.4), Inches(1.2), SOFT)
        add_text(s, name, Inches(x + 0.15), Inches(5.92), Inches(2.1), Inches(0.32), 13, MUTED, True, align=PP_ALIGN.CENTER)
        add_text(s, "localhost" + port, Inches(x + 0.1), Inches(6.28), Inches(2.2), Inches(0.4), 16, NAVY, True, align=PP_ALIGN.CENTER)
        x += 2.55
    footer(s, 4)
    notes(s, "Si hay demo, esta diapositiva es el mapa. DAG real: weekly_airline_retraining. Tareas: monitor_raw_data, run_training_pipeline, validate_champion, reload_api_model, validate_api. Decir que el modelo se recarga en caliente.")

    # ----- 5. Contrato -----
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    header(s, "El contrato de la API", "POST /predict  ·  modelo airline_satisfaction_without_scores")
    add_round(s, Inches(0.45), Inches(1.5), Inches(6.15), Inches(3.85), WHITE, LINE)
    add_text(s, "ENTRADA", Inches(0.7), Inches(1.68), Inches(5.6), Inches(0.28), 12, TEAL, True)
    req = [
        "Gender, Customer Type, Type of Travel, Class",
        "Age entre 0 y 120",
        "Flight Distance mayor o igual a 0",
        "Departure Delay y Arrival Delay mayores o iguales a 0",
        "Cualquier campo extra está prohibido",
    ]
    write_lines(s, req, Inches(0.7), Inches(2.15), Inches(5.6), Inches(2.9), size=15, color=INK, spacing=8)

    add_round(s, Inches(6.8), Inches(1.5), Inches(6.05), Inches(3.85), WHITE, LINE)
    add_text(s, "SALIDA", Inches(7.05), Inches(1.68), Inches(5.5), Inches(0.28), 12, TEAL, True)
    res = [
        "prediction y label",
        "probability",
        "model_name",
        "model_alias  ·  champion",
        "model_version",
    ]
    write_lines(s, res, Inches(7.05), Inches(2.15), Inches(5.5), Inches(2.9), size=15, color=INK, spacing=8)

    add_round(s, Inches(0.45), Inches(5.5), Inches(6.15), Inches(1.5), WHITE, LINE)
    add_text(s, "422", Inches(0.7), Inches(5.68), Inches(1.1), Inches(0.36), 20, GOLD, True)
    add_text(s, "Edad fuera de rango, demora negativa o un campo de más.", Inches(0.7), Inches(6.15), Inches(5.6), Inches(0.6), 14, INK)

    add_round(s, Inches(6.8), Inches(5.5), Inches(6.05), Inches(1.5), WHITE, LINE)
    add_text(s, "503", Inches(7.05), Inches(5.68), Inches(1.1), Inches(0.36), 20, GOLD, True)
    add_text(s, "Sin modelo en memoria, la predicción responde 503. /health sigue disponible.", Inches(7.05), Inches(6.15), Inches(5.5), Inches(0.6), 14, INK)
    footer(s, 5)
    notes(s, "En la demo: un JSON válido en Swagger y uno con Age 200 para ver el 422. /health y POST /reload-model forman parte del servicio: Airflow llama a reload cuando termina el reentrenamiento.")

    # ----- 6. Gobernanza -----
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    header(s, "Seguridad y gobernanza", "La protección recorre datos, aplicación, modelo y operación.")
    cards = [
        ("Datos", "Airflow valida el esquema, los nulos bajo el 10 %, la edad y las demoras. Si el control falla, el entrenamiento no arranca."),
        ("Aplicación", "Pydantic rechaza la entrada inválida con 422. Las credenciales viven en variables de entorno, fuera del repositorio."),
        ("Modelo", "Estabilidad ante ruido, paridad al cambiar Gender y una demora de 300 minutos que no debe subir la satisfacción. SHAP queda en MLflow."),
        ("Operación", "Champion y challenger por ROC AUC. Airflow recarga la API al terminar. La auditoría queda en la corrida robustness_audit."),
    ]
    spots = [(0.45, 1.48), (6.85, 1.48), (0.45, 4.15), (6.85, 4.15)]
    for (title, body), (x, y) in zip(cards, spots):
        add_round(s, Inches(x), Inches(y), Inches(6.05), Inches(2.5), WHITE, LINE)
        add_rect(s, Inches(x), Inches(y), Inches(0.08), Inches(2.5), TEAL)
        add_text(s, title, Inches(x + 0.32), Inches(y + 0.22), Inches(5.4), Inches(0.38), 18, NAVY, True)
        add_text(s, body, Inches(x + 0.32), Inches(y + 0.75), Inches(5.4), Inches(1.5), 15, INK)
    footer(s, 6)
    notes(s, "Nombrar los cuatro componentes y un control de cada uno. Umbrales que se pueden decir: nulos 10 %, estabilidad 90 %, paridad de género 95 %. Si la corrida robustness_audit está en MLflow, mostrar el tag y el gráfico SHAP. No inventar el porcentaje medido.")

    # ----- 7. Demo -----
    s = prs.slides.add_slide(blank)
    set_bg(s, CREAM)
    header(s, "Demo", "Cuatro pantallas del mismo ciclo.")
    steps = [
        ("1", "Streamlit", "localhost:8501", "Un pasajero de clase Business, viaje de negocios, poca demora."),
        ("2", "Swagger", "localhost:8000/docs", "El mismo JSON, y uno con edad 200 para ver el 422."),
        ("3", "MLflow", "localhost:5001", "El experimento, el alias champion y el gráfico SHAP."),
        ("4", "Airflow", "localhost:8080", "Calidad de datos, entrenamiento, champion, recarga y health."),
    ]
    y = 1.5
    for num, title, url, body in steps:
        add_round(s, Inches(0.5), Inches(y), Inches(12.3), Inches(1.3), WHITE, LINE)
        add_round(s, Inches(0.72), Inches(y + 0.36), Inches(0.58), Inches(0.58), TEAL)
        add_text(s, num, Inches(0.72), Inches(y + 0.46), Inches(0.58), Inches(0.4), 18, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, title, Inches(1.55), Inches(y + 0.22), Inches(3.3), Inches(0.4), 18, NAVY, True)
        add_text(s, url, Inches(1.55), Inches(y + 0.64), Inches(3.5), Inches(0.36), 14, TEAL, True)
        add_text(s, body, Inches(5.3), Inches(y + 0.38), Inches(7.1), Inches(0.55), 16, INK, False)
        y += 1.4
    footer(s, 7)
    notes(s, "Orden: Streamlit, Swagger con un caso válido y otro con Age 200, MLflow (champion y SHAP si está la corrida), Airflow DAG weekly_airline_retraining. Si preguntan por otro protocolo, el servicio que se demuestra es esta API. Seguir con la demo.")

    out = "/home/belen/Escritorio/CEIA/MLOPS2/TP_MLOPs2/presentacion/MLOps2_satisfaccion_pasajeros.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
